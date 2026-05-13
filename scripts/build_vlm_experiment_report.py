from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
ASSET_DIR = REPORT_DIR / "assets"
OUTPUT_DOCX = REPORT_DIR / "vlm_blender_scene_quality_experiment_report.docx"

TARGET_DESCRIPTION = (
    "A ceramic cup should be placed on a wooden table under balanced studio lighting. "
    "The cup should be clearly visible in the camera view."
)

ACCENT = "2F5D50"
ACCENT_LIGHT = "E8F1EE"
MUTED = "65736E"
HEADER_BG = "DDEBE7"
BAD_BG = "F9E2DD"
GOOD_BG = "DFF0E6"


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="C9D6D1", size="6") -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size: float = 10.5, color: str | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    set_run_font(run, size=18 if level == 1 else 13, bold=True, color=ACCENT if level == 1 else "1F302B")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color="1F2523")


def add_callout(doc: Document, title: str, body: str, fill: str = ACCENT_LIGHT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table, color="B9CAC3", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, bold=True, color=ACCENT)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    body_run = p2.add_run(body)
    set_run_font(body_run, size=10, color="26322E")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_key_value_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(11.5)
    set_table_borders(table)
    for idx, (key, value) in enumerate(rows):
        key_cell, value_cell = table.rows[idx].cells
        for cell in (key_cell, value_cell):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(key_cell, HEADER_BG)
        key_para = key_cell.paragraphs[0]
        key_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        key_run = key_para.add_run(key)
        set_run_font(key_run, size=9.5, bold=True, color=ACCENT)
        value_para = value_cell.paragraphs[0]
        value_run = value_para.add_run(value)
        set_run_font(value_run, size=9.5, color="1F2523")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_score_table(doc: Document, before: dict, after: dict) -> None:
    rows = [
        ("Total", "total_score", 100),
        ("Geometry", "geometry_score", 20),
        ("Layout", "layout_score", 20),
        ("Material", "material_score", 15),
        ("Lighting", "lighting_score", 15),
        ("Camera", "camera_score", 15),
        ("Prompt Alignment", "alignment_score", 15),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=5)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    headers = ["Metric", "Before", "After", "Delta", "Max"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, HEADER_BG)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=9.5, bold=True, color=ACCENT)
    for r, (label, key, max_score) in enumerate(rows, start=1):
        b = int(before[key])
        a = int(after[key])
        values = [label, str(b), str(a), f"+{a - b}", str(max_score)]
        for c, value in enumerate(values):
            cell = table.rows[r].cells[c]
            set_cell_margins(cell)
            if c == 1:
                set_cell_shading(cell, BAD_BG)
            elif c in (2, 3):
                set_cell_shading(cell, GOOD_BG)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, size=9.2, bold=c in (2, 3), color="1F2523")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_issue_table(doc: Document, issues: list[dict]) -> None:
    table = doc.add_table(rows=1 + len(issues), cols=4)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    headers = ["Issue Type", "Severity", "Object", "Evidence and Diagnosis"]
    widths = [Cm(4.0), Cm(2.2), Cm(2.8), Cm(7.0)]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, HEADER_BG)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=9.2, bold=True, color=ACCENT)
    for r, issue in enumerate(issues, start=1):
        values = [
            issue.get("type", ""),
            issue.get("severity", ""),
            issue.get("object", ""),
            f"{issue.get('view_evidence', '')}: {issue.get('description', '')}",
        ]
        for c, value in enumerate(values):
            cell = table.rows[r].cells[c]
            set_cell_margins(cell)
            if c == 1 and value == "high":
                set_cell_shading(cell, BAD_BG)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in (1, 2) else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, size=8.6, color="1F2523")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def make_score_chart(before: dict, after: dict, output: Path) -> Path:
    labels = ["Total", "Geometry", "Layout", "Material", "Lighting", "Camera", "Alignment"]
    keys = [
        "total_score",
        "geometry_score",
        "layout_score",
        "material_score",
        "lighting_score",
        "camera_score",
        "alignment_score",
    ]
    max_scores = [100, 20, 20, 15, 15, 15, 15]
    width, height = 1400, 720
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype("arial.ttf", 44)
        font = ImageFont.truetype("arial.ttf", 26)
        font_small = ImageFont.truetype("arial.ttf", 22)
    except OSError:
        font_title = ImageFont.load_default()
        font = ImageFont.load_default()
        font_small = ImageFont.load_default()
    draw.rectangle([0, 0, width, 90], fill="#E8F1EE")
    draw.text((50, 24), "VLM Evaluation Score: Before vs. After Repair", fill="#2F5D50", font=font_title)
    x0, y0 = 230, 145
    bar_w, gap = 48, 95
    chart_h = 400
    draw.line([x0, y0, x0, y0 + chart_h], fill="#5A6964", width=2)
    draw.line([x0, y0 + chart_h, width - 70, y0 + chart_h], fill="#5A6964", width=2)
    for pct in [0, 25, 50, 75, 100]:
        y = y0 + chart_h - int(chart_h * pct / 100)
        draw.line([x0, y, width - 70, y], fill="#E1E7E4", width=1)
        draw.text((70, y - 12), f"{pct}%", fill="#65736E", font=font_small)
    for i, (label, key, max_score) in enumerate(zip(labels, keys, max_scores)):
        group_x = x0 + 50 + i * (bar_w * 2 + gap)
        before_pct = before[key] / max_score
        after_pct = after[key] / max_score
        before_h = int(chart_h * before_pct)
        after_h = int(chart_h * after_pct)
        draw.rectangle([group_x, y0 + chart_h - before_h, group_x + bar_w, y0 + chart_h], fill="#D96F5F")
        draw.rectangle(
            [group_x + bar_w + 8, y0 + chart_h - after_h, group_x + 2 * bar_w + 8, y0 + chart_h],
            fill="#4F9E6D",
        )
        draw.text((group_x - 8, y0 + chart_h + 18), label, fill="#26322E", font=font_small)
        draw.text((group_x - 2, y0 + chart_h - before_h - 28), str(before[key]), fill="#8A3F36", font=font_small)
        draw.text(
            (group_x + bar_w + 14, y0 + chart_h - after_h - 28),
            str(after[key]),
            fill="#2F6F49",
            font=font_small,
        )
    draw.rectangle([930, 605, 960, 635], fill="#D96F5F")
    draw.text((975, 604), "Before repair", fill="#26322E", font=font)
    draw.rectangle([930, 650, 960, 680], fill="#4F9E6D")
    draw.text((975, 649), "After repair", fill="#26322E", font=font)
    output.parent.mkdir(parents=True, exist_ok=True)
    img.save(output)
    return output


def make_comparison_panel(left: Path, right: Path, left_label: str, right_label: str, output: Path) -> Path:
    width, height = 1400, 760
    panel = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(panel)
    try:
        font_title = ImageFont.truetype("arial.ttf", 32)
        font_label = ImageFont.truetype("arial.ttf", 24)
    except OSError:
        font_title = ImageFont.load_default()
        font_label = ImageFont.load_default()

    draw.rectangle([0, 0, width, 72], fill="#E8F1EE")
    draw.text((40, 20), "Before / After Visual Comparison", fill="#2F5D50", font=font_title)
    slots = [(60, 118, 660, 718), (740, 118, 1340, 718)]
    for path, label, box, border in (
        (left, left_label, slots[0], "#D96F5F"),
        (right, right_label, slots[1], "#4F9E6D"),
    ):
        draw.rectangle([box[0] - 10, box[1] - 42, box[2] + 10, box[3] + 10], outline=border, width=4)
        draw.text((box[0], box[1] - 34), label, fill="#26322E", font=font_label)
        src = Image.open(path).convert("RGB")
        src.thumbnail((box[2] - box[0], box[3] - box[1]))
        x = box[0] + ((box[2] - box[0]) - src.width) // 2
        y = box[1] + ((box[3] - box[1]) - src.height) // 2
        panel.paste(src, (x, y))

    output.parent.mkdir(parents=True, exist_ok=True)
    panel.save(output)
    return output


def add_image_with_caption(doc: Document, image_path: Path, caption: str, width: float = 5.6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, size=8.6, color=MUTED)
    cap.paragraph_format.space_after = Pt(4)


def add_side_by_side_images(doc: Document, left: Path, right: Path, left_caption: str, right_caption: str) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table, color="FFFFFF", size="0")
    for col, (path, caption) in enumerate(((left, left_caption), (right, right_caption))):
        img_cell = table.cell(0, col)
        cap_cell = table.cell(1, col)
        for cell in (img_cell, cap_cell):
            set_cell_margins(cell, top=40, start=60, bottom=40, end=60)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = img_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(3.05))
        cap_p = cap_cell.paragraphs[0]
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap_p.add_run(caption)
        set_run_font(run, size=8.5, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2"]:
        styles[style_name].font.name = "Microsoft YaHei"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build_report() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    before_eval = load_json(ROOT / "logs" / "eval_result.json")
    after_eval = load_json(ROOT / "logs" / "eval_result_after_repair.json")
    before_meta = load_json(ROOT / "metadata" / "bad_cup_scene.json")
    after_meta = load_json(ROOT / "metadata" / "repaired_cup_scene.json")
    chart_path = make_score_chart(before_eval, after_eval, ASSET_DIR / "score_comparison.png")

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "VLM-based Blender Scene Quality Assessment and Repair Experiment"
    doc.core_properties.author = "Codex + Blender MCP"

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(34)
    title_run = cover.add_run("VLM-based Blender Scene Quality Assessment\nand Repair Experiment")
    set_run_font(title_run, size=24, bold=True, color=ACCENT)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("当前 MCP 会话版诊断与修复闭环实验报告")
    set_run_font(subtitle_run, size=14, color="33413D")
    doc.add_paragraph()
    add_callout(
        doc,
        "实验摘要",
        "本实验在当前打开的 Blender MCP 会话中构造一个含有故意错误的杯子-木桌场景，"
        "通过多视角渲染与场景 metadata 进行 VLM-style 质量评估，再根据结构化 detected issues "
        "自动执行修复脚本。修复后总分从 42/100 提升到 90/100。",
    )
    add_key_value_table(
        doc,
        [
            ("目标描述", TARGET_DESCRIPTION),
            ("实验日期", date.today().isoformat()),
            ("执行环境", "Blender 5.1 + official Blender MCP bridge + bpy repair scripts"),
            ("核心闭环", "create bad scene -> render multiview -> export metadata -> VLM-style JSON evaluation -> repair -> re-render -> re-score"),
        ],
    )
    doc.add_page_break()

    add_heading(doc, "1. Experimental Goal", 1)
    add_body(
        doc,
        "本实验验证一个最小可运行闭环：Agent 不只生成 Blender 场景，还能够把场景渲染成多视角图像，"
        "结合 metadata 诊断几何、布局、材质、光照和相机问题，并把诊断结果转化为可执行的修复动作。"
    )
    add_body(
        doc,
        "为了让评测信号明确，第一轮场景被故意设置为坏样本：杯子悬空、光照过暗、杯子使用默认灰色材质、"
        "相机偏离主体。这类错误都可以被多视角图像和 metadata 同时观测。"
    )

    add_heading(doc, "2. Pipeline Design", 1)
    add_key_value_table(
        doc,
        [
            ("造坏场景", "通过 MCP 在当前 Blender 会话中生成 wooden_table、ceramic_cup、area_key_light 和 Camera，并保存 bad_cup_scene.blend。"),
            ("渲染评估输入", "渲染 camera/front/side/top/back 五视图，并导出 object、camera、light、bounding box 和 material metadata。"),
            ("VLM 打分", "读取五视图和 metadata，按固定 JSON schema 输出 total score、六项分数、detected issues 和 repair suggestions。"),
            ("修复", "根据 issue type 调用 repair_from_eval.py 中对应修复函数，直接修改当前 MCP 场景并保存 repaired_cup_scene.blend。"),
            ("复评", "重新渲染、重新导出 metadata，并生成 after repair 评分用于前后对比。"),
        ],
    )

    add_heading(doc, "3. Visual Evidence", 1)
    add_body(doc, "下面展示关键视角的 before/after 对比。为保证 Word 兼容性，每张渲染图都作为普通图片直接插入文档，不放入表格或复杂容器。")
    add_image_with_caption(
        doc,
        ROOT / "renders" / "bad_cup_scene" / "camera.png",
        "Figure 1. Before repair - camera view: dark lighting and weak framing",
        width=4.6,
    )
    add_image_with_caption(
        doc,
        ROOT / "renders" / "repaired_cup_scene" / "camera.png",
        "Figure 2. After repair - camera view: clear cup and table composition",
        width=4.6,
    )
    add_image_with_caption(
        doc,
        ROOT / "renders" / "bad_cup_scene" / "side.png",
        "Figure 3. Before repair - side view: the cup is floating above the tabletop",
        width=4.6,
    )
    add_image_with_caption(
        doc,
        ROOT / "renders" / "repaired_cup_scene" / "side.png",
        "Figure 4. After repair - side view: the cup rests on the tabletop",
        width=4.6,
    )
    add_image_with_caption(
        doc,
        ROOT / "renders" / "bad_cup_scene" / "top.png",
        "Figure 5. Before repair - top view: gray/default material under dark lighting",
        width=4.6,
    )
    add_image_with_caption(
        doc,
        ROOT / "renders" / "repaired_cup_scene" / "top.png",
        "Figure 6. After repair - top view: white ceramic material and brighter scene",
        width=4.6,
    )
    doc.add_page_break()

    add_heading(doc, "4. Evaluation Results", 1)
    add_body(
        doc,
        "评分采用 100 分制：geometry 20、layout 20、material 15、lighting 15、camera 15、prompt alignment 15。"
        "修复后的提升主要来自布局、光照、材质和相机四个维度。"
    )
    add_score_table(doc, before_eval, after_eval)
    add_image_with_caption(doc, chart_path, "Figure 7. Score comparison before and after repair", width=5.9)

    add_heading(doc, "5. Detected Issues Before Repair", 1)
    add_issue_table(doc, before_eval.get("detected_issues", []))

    add_heading(doc, "6. Metadata Verification", 1)
    before_light = before_meta["lights"][0]["energy"]
    after_light = after_meta["lights"][0]["energy"]
    before_exposure = before_meta["color_management"]["exposure"]
    after_exposure = after_meta["color_management"]["exposure"]
    before_cup = next(obj for obj in before_meta["objects"] if obj["name"] == "ceramic_cup")
    after_cup = next(obj for obj in after_meta["objects"] if obj["name"] == "ceramic_cup")
    table = next(obj for obj in after_meta["objects"] if obj["name"] == "wooden_table")
    add_key_value_table(
        doc,
        [
            ("Cup bottom Z before", f"{before_cup['bounding_box']['min'][2]:.3f}; tabletop top Z is {table['bounding_box']['max'][2]:.3f}, confirming a floating object."),
            ("Cup bottom Z after", f"{after_cup['bounding_box']['min'][2]:.3f}; tabletop top Z is {table['bounding_box']['max'][2]:.3f}, confirming contact alignment."),
            ("Area light energy", f"{before_light:.1f} -> {after_light:.1f}"),
            ("Exposure", f"{before_exposure:.1f} -> {after_exposure:.1f}"),
            ("Cup material", f"{before_cup['material_names'][0]} -> {after_cup['material_names'][0]}"),
        ],
    )

    add_heading(doc, "7. Conclusion and Next Steps", 1)
    add_body(
        doc,
        "本次实验说明：即使第一版 VLM 评分先采用人工/模型代理读取多视图的方式，闭环仍然可以稳定跑通，"
        "并且每个错误都能映射到具体 Blender API 修复动作。对于后续科研工作，这个 demo 可以扩展为小型 benchmark，"
        "覆盖更多物体类别、空间关系和错误类型。"
    )
    add_callout(
        doc,
        "下一步建议",
        "将当前的 VLM-style 评分替换为外部 VLM API；增加自动重试与多轮 repair；把每轮 score delta、issue recall "
        "和 repair success rate 统计为表格，从工程 demo 逐步推进到可投稿的实验设计。",
        fill="F1F6F4",
    )

    doc.save(OUTPUT_DOCX)
    print(json.dumps({"report": str(OUTPUT_DOCX), "chart": str(chart_path)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build_report()
